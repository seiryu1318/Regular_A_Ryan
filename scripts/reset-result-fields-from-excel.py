from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "public" / "admissions-data.json"
CORRECTIONS_PATH = ROOT / "audit" / "official-results-corrections.json"
EXCEL_PATH = Path(
    r"C:\Users\User\Desktop\업무\연도\2026년\03. 개인\05. 개인 상담자료\여고 상담\03. 참고 자료\대학어디가(2023~2026)정시입결자료(남악고 김현석)ver0702.xlsx"
)
FIELD_COLUMNS = {"n": 8, "c": 9, "add": 10, "cv50": 11, "cv70": 12, "max": 13, "p50": 14, "p70": 15}


def main() -> None:
    payload = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    workbook = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    sheet = workbook["대학자료"]
    excel_rows = {row_number - 1: row for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2)}

    restored = 0
    for row in payload["scores"]:
        source = excel_rows[int(row["id"])]
        for field, column in FIELD_COLUMNS.items():
            row[field] = source[column]
            restored += 1

    DATA_PATH.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    CORRECTIONS_PATH.write_text("[]", encoding="utf-8")
    print(json.dumps({"rows": len(payload["scores"]), "restored_fields": restored}, ensure_ascii=False))


if __name__ == "__main__":
    main()
