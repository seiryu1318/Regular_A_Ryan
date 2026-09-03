from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "public" / "admissions-data.json"
CHANGE_PATH = ROOT / "audit" / "official-results-corrections.json"
CONSISTENCY_CHANGE_PATH = ROOT / "audit" / "consistency-corrections.json"
PLACEHOLDER_CHANGE_PATH = ROOT / "audit" / "placeholder-value-corrections.json"
TRACK_CHANGE_PATH = ROOT / "audit" / "track-corrections.json"
FINAL_CONSISTENCY_CHANGE_PATH = ROOT / "audit" / "final-consistency-corrections.json"
SUMMARY_PATH = ROOT / "audit" / "excel-source-summary.json"
DIFFERENCES_PATH = ROOT / "audit" / "excel-source-differences.json"
EXCEL_PATH = Path(
    r"C:\Users\User\Desktop\업무\연도\2026년\03. 개인\05. 개인 상담자료\여고 상담\03. 참고 자료\대학어디가(2023~2026)정시입결자료(남악고 김현석)ver0702.xlsx"
)
FIELD_COLUMNS = {
    "u": 1,
    "y": 2,
    "a": 5,
    "d": 6,
    "t": 7,
    "n": 8,
    "c": 9,
    "add": 10,
    "cv50": 11,
    "cv70": 12,
    "max": 13,
    "p50": 14,
    "p70": 15,
    "ko": 16,
    "ma": 17,
    "inq": 18,
    "en": 20,
    "x": 23,
}


def equal(left, right) -> bool:
    if isinstance(left, str) and not left.strip():
        left = None
    if isinstance(right, str) and not right.strip():
        right = None
    if left is None and right in (None, ""):
        return True
    if right is None and left in (None, ""):
        return True
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return abs(float(left) - float(right)) <= 0.011
    return str(left).strip() == str(right).strip()


def clean_department(value) -> str:
    return "".join(character for character in str(value or "") if character.isalnum() and character not in "ㆍ")


data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
corrections = json.loads(CHANGE_PATH.read_text(encoding="utf-8")) if CHANGE_PATH.exists() else []
consistency_corrections = json.loads(CONSISTENCY_CHANGE_PATH.read_text(encoding="utf-8")) if CONSISTENCY_CHANGE_PATH.exists() else []
placeholder_corrections = json.loads(PLACEHOLDER_CHANGE_PATH.read_text(encoding="utf-8")) if PLACEHOLDER_CHANGE_PATH.exists() else []
track_corrections = json.loads(TRACK_CHANGE_PATH.read_text(encoding="utf-8")) if TRACK_CHANGE_PATH.exists() else []
final_consistency_corrections = json.loads(FINAL_CONSISTENCY_CHANGE_PATH.read_text(encoding="utf-8")) if FINAL_CONSISTENCY_CHANGE_PATH.exists() else []
correction_keys = {
    (int(item["id"]), item["field"])
    for item in corrections + consistency_corrections + placeholder_corrections + track_corrections + final_consistency_corrections
}
json_rows = {int(row["id"]): row for row in data["scores"]}

workbook = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
sheet = workbook["대학자료"]
excel_rows = {}
for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
    if row[3] != "정시" or row[2] not in (2024, 2025, 2026):
        continue
    excel_rows[row_number - 1] = row

differences = []
for row_id, excel_row in excel_rows.items():
    json_row = json_rows.get(row_id)
    if not json_row:
        continue
    for field, column in FIELD_COLUMNS.items():
        if field == "d" and clean_department(excel_row[column]) == clean_department(json_row.get(field)):
            continue
        if equal(excel_row[column], json_row.get(field)):
            continue
        differences.append(
            {
                "id": row_id,
                "field": field,
                "excel": excel_row[column],
                "json": json_row.get(field),
                "official_correction": (row_id, field) in correction_keys or field == "x",
            }
        )

summary = {
    "excel_rows": len(excel_rows),
    "json_rows": len(json_rows),
    "missing_json_rows": len(set(excel_rows) - set(json_rows)),
    "extra_json_rows": len(set(json_rows) - set(excel_rows)),
    "field_differences": len(differences),
    "official_or_recomputed_differences": sum(item["official_correction"] for item in differences),
    "unexplained_differences": sum(not item["official_correction"] for item in differences),
    "unexplained_by_field": {
        field: sum(item["field"] == field and not item["official_correction"] for item in differences)
        for field in FIELD_COLUMNS
        if any(item["field"] == field and not item["official_correction"] for item in differences)
    },
}
SUMMARY_PATH.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
DIFFERENCES_PATH.write_text(json.dumps(differences, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=2))
