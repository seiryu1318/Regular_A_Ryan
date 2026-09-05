from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "public" / "admissions-data.json"
SOURCE_PATH = ROOT / "audit" / "official-results-files" / "gachon-2026-results.xlsx"
AUDIT_PATH = ROOT / "audit" / "gachon-2026-refresh.json"
SOURCE_URL = "https://admission.gachon.ac.kr/admission/html/regular/resultView.asp?BOARD_IDX=32175"
ADMISSION_URL = "https://admission.gachon.ac.kr/admission/html/regular/main.asp"

SPECIAL_ADMISSION = re.compile(
    r"특성화|농어촌|지역(?:인재|균형|기회|메디|전형)|강원인재|기회균형|기회균등|"
    r"저소득|차상위|수급|사회(?:적)?배려|특수교육|장애인|성인학습|재직자|군위탁|"
    r"정원외|서해5도|고른기회|계약학과|기독교전형|군사학과전형|국방.*전형|"
    r"사이버국방|항공시스템공학.*특별|자율전공\s*특별"
)

TRACK_OVERRIDES = {
    "AI인문대학": "인문",
    "경제학과": "인문",
    "관광경영학과": "인문",
    "금융·빅데이터학부": "인문",
    "도시계획·조경학부": "자연",
    "법과대학": "인문",
    "사회복지학과": "인문",
    "심리학과": "인문",
    "유아교육학과": "인문",
    "회계세무학과": "인문",
    "경영학과": "인문",
    "미디어커뮤니케이션학과": "인문",
    "의료산업경영학과": "인문",
    "패션산업학과": "인문",
    "의예과": "의약학",
    "한의예과": "의약학",
    "약학과": "의약학",
    "간호학과": "자연",
    "물리치료학과": "자연",
    "방사선학과": "자연",
    "응급구조학과": "자연",
    "치위생학과": "자연",
    "연기예술학과(연출)": "예체능",
    "운동재활학과": "예체능",
}


def clean_name(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z가-힣]", "", value or "")


def infer_track(department: str, existing_tracks: dict[str, str]) -> str:
    if department in TRACK_OVERRIDES:
        return TRACK_OVERRIDES[department]
    known = existing_tracks.get(clean_name(department))
    if known:
        return known
    if re.search(r"간호|치위생|방사선|물리치료|응급구조|임상병리|작업치료", department):
        return "자연"
    if re.search(r"의예|한의예|치의예|약학|수의", department):
        return "의약학"
    if re.search(r"연기|음악|미술|디자인|체육|운동재활", department):
        return "예체능"
    if re.search(r"경제|경영|법과|사회|심리|교육|회계|관광|미디어|패션|인문", department):
        return "인문"
    return "자연"


def main() -> None:
    payload = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    original_scores = payload["scores"]

    removed_special = [
        row for row in original_scores
        if SPECIAL_ADMISSION.search(f"{row.get('exam') or ''} {row.get('a') or ''}")
    ]
    retained = [
        row for row in original_scores
        if row not in removed_special and not (row.get("u") == "가천대" and row.get("y") == 2026)
    ]

    existing_tracks: dict[str, str] = {}
    for row in original_scores:
        if row.get("u") != "가천대" or not row.get("t"):
            continue
        existing_tracks.setdefault(clean_name(row.get("d", "")), row["t"])

    workbook = load_workbook(SOURCE_PATH, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    next_id = max((int(row.get("id", 0)) for row in retained), default=0) + 1
    added: list[dict] = []

    for values in sheet.iter_rows(min_row=4, values_only=True):
        group, exam, department, headcount, competition, score50, score70, _score90, _waitlist = values[:9]
        if not group or not exam or not department:
            continue
        exam_name = str(exam).strip()
        if exam_name not in {"일반", "일반2"}:
            continue
        group_letter = str(group).strip()[0]
        metric = "표준점수" if exam_name == "일반2" else "백분위"
        row = {
            "id": next_id,
            "u": "가천대",
            "y": 2026,
            "a": f"정시({group_letter})",
            "g": group_letter,
            "d": str(department).strip(),
            "t": infer_track(str(department).strip(), existing_tracks),
            "n": int(headcount) if isinstance(headcount, (int, float)) else None,
            "c": float(competition) if isinstance(competition, (int, float)) else None,
            "x": None,
            "add": None,
            "cv50": None,
            "cv70": None,
            "max": None,
            "p50": float(score50) if isinstance(score50, (int, float)) else None,
            "p70": float(score70) if isinstance(score70, (int, float)) else None,
            "ko": None,
            "ma": None,
            "inq": None,
            "en": None,
            "r": "경기",
            "source": SOURCE_URL,
            "admission": ADMISSION_URL,
            "metric": metric,
            "sb": None,
            "exam": f"수능({exam_name}전형)",
        }
        added.append(row)
        next_id += 1

    payload["scores"] = retained + added
    DATA_PATH.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    AUDIT_PATH.write_text(
        json.dumps(
            {
                "source": SOURCE_URL,
                "source_file": str(SOURCE_PATH),
                "removed_special_rows": len(removed_special),
                "removed_special_by_exam": sorted({row.get("exam") or "" for row in removed_special}),
                "added_gachon_2026_rows": len(added),
                "added_general_rows": sum(row["metric"] == "백분위" for row in added),
                "added_general2_rows": sum(row["metric"] == "표준점수" for row in added),
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "scores_before": len(original_scores),
                "special_removed": len(removed_special),
                "gachon_2026_added": len(added),
                "scores_after": len(payload["scores"]),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
